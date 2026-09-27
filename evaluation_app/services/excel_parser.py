"""Парсинг Excel-файлов преподавателей (openpyxl).

Поддерживаются два реальных формата оценочных листов:

1. **Строки** (шаблон из плана): ``ФИО | Название конкурса | Итого`` —
   одна строка = одна оценка ученика за конкурс.
2. **Критерии** (оценочная матрица): ``ФИО | критерий 1 | ... | критерий N | ИТОГО``,
   баллы начисляются в ячейках критериев. Служебные столбцы («№», «класс» и т. п.)
   критериями не считаются; баллы служебного столбца «за вход в систему» входят
   в итог ученика отдельной оценкой, если формула ИТОГО их не учитывает.

Шапка реальных листов многоуровневая: группы направлений, названия критериев и
строка максимумов по критериям. Поэтому структура определяется по блоку из 1–6 строк,
а строки со значениями (нумерация столбцов, максимумы) при выборе подписи пропускаются.
"""

from __future__ import annotations

import contextlib
import os
import zipfile
from collections.abc import Iterable, Sequence
from dataclasses import dataclass, field
from typing import Any

import openpyxl

from evaluation_app import config
from evaluation_app.models import Direction, Evaluation, Teacher
from evaluation_app.utils.file_utils import safe_file_stem
from evaluation_app.utils.validators import (
    DataValidator,
    detect_columns,
    exact_match_column,
    is_header_or_total_marker,
    is_login_column,
    is_service_label,
    is_total_marker,
    match_column,
    missing_fields,
    name_key,
    normalize_header,
    normalize_name,
    parse_score,
    shorten,
)

MODE_ROWS = "rows"
MODE_CRITERIA = "criteria"

# Сколько строк может занимать блок заголовков.
MAX_HEADER_BLOCK_ROWS = 6
# Максимум предупреждений по одному листу.
MAX_SHEET_WARNINGS = 10
# Сколько строк без ФИО подряд означают конец данных («фантомные» строки листа).
MAX_EMPTY_ROWS_RUN = 100
# Длинные подписи (описания критериев) не могут быть столбцом с ФИО.
MAX_STUDENT_LABEL_LENGTH = 32

# Подписи листов-сводов: их данные обычно уже собраны из файлов преподавателей.
SUMMARY_SHEET_MARKERS = ("общий свод", "свод", "сводная", "итоговый", "общая таблица")


class ExcelParseError(Exception):
    """Ошибка чтения или разбора Excel-файла."""


@dataclass
class ColumnInfo:
    """Сведения о столбце, собранные из блока заголовков."""

    index: int
    label: str = ""  # подпись для поиска поля и названия критерия
    group_label: str = ""  # подпись группы (верхний уровень)
    display_name: str = ""  # имя для интерфейса
    field: str = ""  # student_name / competition_name / total_score


@dataclass
class HeaderLayout:
    """Результат определения структуры листа."""

    header_rows: tuple[int, int]
    columns: dict[int, ColumnInfo]
    student_column: int
    total_column: int | None = None
    competition_column: int | None = None
    criterion_columns: tuple[int, ...] = ()
    # Служебный столбец с баллами входа («за вход в систему»): формула ИТОГО
    # реальных файлов его не учитывает, поэтому его баллы прибавляются к итогу.
    login_column: int | None = None
    mode: str = MODE_ROWS

    @property
    def first_header_row(self) -> int:
        """Первая строка блока заголовков."""
        return self.header_rows[0]

    @property
    def first_data_row(self) -> int:
        """Первая строка с данными."""
        return self.header_rows[1] + 1

    def display_name(self, column: int) -> str:
        """Имя столбца для интерфейса."""
        info = self.columns.get(column)
        if info is None:
            return f"Столбец {column}"
        return info.display_name or info.label or f"Столбец {column}"


@dataclass
class SheetParseResult:
    """Результат разбора одного листа Excel."""

    sheet_name: str
    direction: Direction
    header_row: int = 0
    header_rows: tuple[int, int] = (0, 0)
    columns: dict[str, int] = field(default_factory=dict)
    mode: str = MODE_ROWS
    criterion_columns: tuple[int, ...] = ()
    criterion_names: tuple[str, ...] = ()  # заполняется при разборе, читается только тестами
    skipped_rows: list[int] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    rows_parsed: int = 0

    @property
    def evaluations_count(self) -> int:
        """Количество оценок на листе."""
        return len(self.direction.evaluations)

    @property
    def total_score(self) -> float:
        """Сумма баллов листа."""
        return self.direction.total_score()

    def to_rows(self, limit: int | None = None) -> list[list[str]]:
        """Строки для предпросмотра в интерфейсе."""
        rows: list[list[str]] = []
        for evaluation in self.direction.evaluations:
            rows.append(
                [
                    evaluation.student_name,
                    evaluation.competition_name,
                    f"{evaluation.total_score:g}",
                ]
            )
            if limit is not None and len(rows) >= limit:
                break
        return rows

    def describe_columns(self) -> str:
        """Текстовая расшифровка структуры листа для диалога импорта."""
        from openpyxl.utils import get_column_letter

        parts: list[str] = []
        if "student_name" in self.columns:
            parts.append(f"ФИО → {get_column_letter(self.columns['student_name'])}")
        if "competition_name" in self.columns:
            parts.append(f"конкурс → {get_column_letter(self.columns['competition_name'])}")
        if "total_score" in self.columns:
            parts.append(f"итого → {get_column_letter(self.columns['total_score'])}")
        if self.mode == MODE_CRITERIA:
            parts.append(f"критериев: {len(self.criterion_columns)}")
        if self.header_rows[1] > self.header_rows[0]:
            parts.append(f"заголовок: строки {self.header_rows[0]}–{self.header_rows[1]}")
        return ", ".join(parts) if parts else "структура не определена"


class ExcelParser:
    """Читает файл преподавателя: каждый лист — направление."""

    def __init__(
        self,
        file_path: str,
        teacher_name: str | None = None,
        header_search_limit: int = 12,
        read_only: bool = True,
    ) -> None:
        self.file_path = str(file_path)
        self.teacher_name = normalize_name(teacher_name) if teacher_name else ""
        self.header_search_limit = max(1, int(header_search_limit))
        self.read_only = read_only
        self.workbook: Any | None = None
        self.header_row_override: int | None = None
        self.column_overrides: dict[str, int] = {}
        self.mode_override: str | None = None
        self.external_links: list[str] = []
        self.warnings: list[str] = []

    # --- Загрузка ------------------------------------------------------------

    def load(self, data_only: bool = True) -> ExcelParser:
        """Открывает Excel-файл."""
        if not os.path.isfile(self.file_path):
            raise ExcelParseError(f"Файл не найден: {self.file_path}")

        extension = os.path.splitext(self.file_path)[1].lower()
        if extension == ".xls":
            raise ExcelParseError(
                "Формат .xls (Excel 97–2003) не поддерживается. "
                "Откройте файл в Excel и сохраните его как .xlsx."
            )
        if extension not in config.EXCEL_EXTENSIONS:
            raise ExcelParseError(
                f"Неподдерживаемый формат файла «{extension or 'без расширения'}». "
                f"Ожидается: {', '.join(config.EXCEL_EXTENSIONS)}"
            )

        try:
            self.workbook = openpyxl.load_workbook(
                self.file_path, data_only=data_only, read_only=self.read_only
            )
        except Exception as exc:
            raise ExcelParseError(f"Не удалось открыть файл: {exc}") from exc

        self.external_links = self._detect_external_links()
        return self

    def _detect_external_links(self) -> list[str]:
        """Ищет ссылки на другие книги (признак сводного файла)."""
        if not zipfile.is_zipfile(self.file_path):
            return []
        try:
            with zipfile.ZipFile(self.file_path) as archive:
                return [name for name in archive.namelist() if name.startswith("xl/externalLinks/")]
        except (OSError, zipfile.BadZipFile):  # pragma: no cover - защита от битых файлов
            return []

    def close(self) -> None:
        """Закрывает книгу (освобождает файл)."""
        if self.workbook is not None:
            with contextlib.suppress(Exception):  # pragma: no cover - закрытие не критично
                self.workbook.close()
            self.workbook = None

    def __enter__(self) -> ExcelParser:
        return self.load()

    def __exit__(self, *_exc_info: Any) -> None:
        self.close()

    def _ensure_workbook(self) -> Any:
        if self.workbook is None:
            self.load()
        return self.workbook

    # --- Вспомогательные методы ---------------------------------------------

    @property
    def sheet_names(self) -> list[str]:
        """Список листов книги."""
        return list(self._ensure_workbook().sheetnames)

    def worksheet(self, sheet_name: str | None = None) -> Any:
        """Лист по имени (по умолчанию — первый)."""
        workbook = self._ensure_workbook()
        if sheet_name is None:
            return workbook[workbook.sheetnames[0]]
        if sheet_name not in workbook.sheetnames:
            raise ExcelParseError(f"Лист «{sheet_name}» не найден в файле")
        return workbook[sheet_name]

    def default_teacher_name(self) -> str:
        """ФИО преподавателя по умолчанию — из имени файла."""
        return self.teacher_name or safe_file_stem(self.file_path)

    def set_overrides(
        self,
        header_row: int | None = None,
        columns: dict[str, int] | None = None,
        mode: str | None = None,
    ) -> None:
        """Задаёт вручную строку заголовков, столбцы и режим разбора листа.

        Значения учитываются при построении структуры листа
        (:meth:`_layout_from_overrides`) и действуют на последующие вызовы
        :meth:`parse_sheet`. Используется только в тестах: диалог ручного
        определения столбцов передаёт разметку прямо в ``parse_sheet``.
        """
        self.header_row_override = header_row
        if columns:
            self.column_overrides = {str(k): int(v) for k, v in columns.items()}
        self.mode_override = mode or self.mode_override

    def _read_rows(self, worksheet: Any, first: int, last: int, max_column: int) -> list[tuple]:
        """Читает строки ``first..last`` (1-based) как кортежи значений."""
        if last < first:
            return []
        rows: list[tuple] = []
        for values in worksheet.iter_rows(
            min_row=first, max_row=last, max_col=max_column, values_only=True
        ):
            rows.append(tuple(values))
        while len(rows) < last - first + 1:
            rows.append(())
        return rows

    def preview_rows(self, sheet_name: str, limit: int = config.SHEET_PREVIEW_ROWS) -> list[list[Any]]:
        """Первые строки листа «как есть» — для ручного определения столбцов."""
        worksheet = self.worksheet(sheet_name)
        max_column = worksheet.max_column or 0
        return [
            list(values)
            for values in self._read_rows(worksheet, 1, min(limit, worksheet.max_row or 0), max_column)
        ]

    @staticmethod
    def _cell(values: Sequence[Any], column: int) -> Any:
        """Значение столбца (1-based) в считанной строке."""
        index = column - 1
        if 0 <= index < len(values):
            return values[index]
        return None

    @staticmethod
    def _digit_ratio(values: Sequence[Any]) -> float:
        """Доля заполненных ячеек строки, содержащих цифры."""
        filled = [value for value in values if normalize_name(value)]
        if not filled:
            return 0.0
        with_digits = sum(1 for value in filled if any(ch.isdigit() for ch in str(value)))
        return with_digits / len(filled)

    @classmethod
    def _is_value_row(cls, values: Sequence[Any]) -> bool:
        """Строка со значениями (нумерация столбцов, максимумы баллов).

        В таких строках почти каждая заполненная ячейка содержит цифры —
        это не подписи критериев, поэтому как названия они не используются.
        """
        filled = [value for value in values if normalize_name(value)]
        if len(filled) < 2:
            return False
        return cls._digit_ratio(values) >= 0.6

    def _header_window(self, worksheet: Any, max_column: int, max_row: int) -> list[tuple]:
        """Первые строки листа одним чтением — шапка разбирается в памяти.

        В режиме read-only каждое обращение к iter_rows читает файл заново,
        поэтому блок заголовков читается один раз.
        """
        last = min(max_row, self.header_search_limit + MAX_HEADER_BLOCK_ROWS + 6)
        return self._read_rows(worksheet, 1, last, max_column)

    def _collect_labels(
        self, window: Sequence[tuple], start: int, size: int, max_column: int
    ) -> dict[int, tuple[str, str, str]]:
        """Подписи столбцов блока: (любая глубокая, из строк-подписей, группа)."""
        rows = [tuple(values) for values in window[start - 1 : start - 1 + size]]
        value_rows = {index for index, values in enumerate(rows) if self._is_value_row(values)}

        labels: dict[int, list[str]] = {}
        for index, values in enumerate(rows):
            for column, value in enumerate(values, start=1):
                text = normalize_name(value)
                if not text:
                    continue
                item = labels.setdefault(column, ["", "", ""])
                item[0] = text
                if index not in value_rows:
                    item[1] = text
                    if not item[2]:
                        item[2] = text
        return {column: (item[0], item[1] or item[0], item[2]) for column, item in labels.items()}

    def _find_header_block(
        self, window: Sequence[tuple], max_column: int, max_row: int
    ) -> tuple[int, int]:
        """Ищет блок строк-заголовков, дающий обязательные столбцы.

        Кандидаты сравниваются по признакам (важные — первыми):
        столбец ФИО найден строго по названию; сколько полей распознано
        (больше — лучше); длина блока (короче — лучше); доля ячеек с цифрами
        (меньше — лучше: в шапке цифр почти нет); позиция (выше — лучше).
        Так шапка не «съезжает» ни на строку данных, где случайно оказались
        слова «ученик» или «баллы», ни на титульную строку отчёта.
        """
        limit = min(self.header_search_limit, max_row)
        # (строгое ФИО, распознано полей, длина блока, номер строки, доля цифр)
        candidates: list[tuple[int, int, int, int, float]] = []

        for start in range(1, limit + 1):
            for size in range(1, MAX_HEADER_BLOCK_ROWS + 1):
                if start + size - 1 > max_row:
                    break
                labels = self._collect_labels(window, start, size, max_column)
                mapping = detect_columns(
                    [labels.get(column, ("", "", ""))[1] for column in range(1, max_column + 1)]
                )
                student_column, exact = self._student_column_from_labels(labels)
                if student_column is None:
                    continue
                # Лист-матрица может не иметь колонки ИТОГО — тогда баллы берутся
                # из столбцов-критериев, и достаточно столбца с ФИО.
                if missing_fields(mapping) and self._criteria_count(labels) < 2:
                    continue
                rows = [tuple(values) for values in window[start - 1 : start - 1 + size]]
                ratios = [self._digit_ratio(values) for values in rows if any(values)]
                digit_ratio = sum(ratios) / len(ratios) if ratios else 0.0
                candidates.append((int(exact), len(mapping), size, start, digit_ratio))

        if not candidates:
            raise ExcelParseError(
                "Не удалось определить структуру таблицы. Убедитесь, что есть столбцы "
                "«ФИО» и «Итого», либо укажите их вручную."
            )

        _exact, _fields, size, start, _digits = max(
            candidates, key=lambda item: (item[0], item[1], -item[2], -item[4], -item[3])
        )
        return start, start + size - 1

    def _extend_header_block(
        self, window: Sequence[tuple], student_column: int, header_rows: tuple[int, int], max_row: int
    ) -> tuple[int, int]:
        """Расширяет блок заголовков, пока в строке нет ФИО ученика.

        В реальных файлах под «шапкой» стоят строки с названиями критериев и
        максимумами баллов — они тоже относятся к заголовку.
        """
        first, last = header_rows
        while last < max_row and (last - first + 1) < MAX_HEADER_BLOCK_ROWS:
            if last >= len(window):
                break
            value = self._cell(window[last], student_column)
            if normalize_name(value):
                break
            last += 1
        return first, last

    def detect_layout(self, worksheet: Any) -> HeaderLayout:
        """Определяет строки заголовков, столбцы полей и режим разбора."""
        max_row = worksheet.max_row or 0
        max_column = worksheet.max_column or 0
        if not max_row or not max_column:
            raise ExcelParseError("Лист пуст — нет данных для разбора")

        window = self._header_window(worksheet, max_column, max_row)
        if self.header_row_override is not None:
            header_rows = (self.header_row_override, self.header_row_override)
        else:
            header_rows = self._find_header_block(window, max_column, max_row)

        columns = self._build_columns(window, header_rows, max_column)

        student_column = self._detect_student_column(columns)
        if student_column is None:
            raise ExcelParseError(
                "Не найден столбец с ФИО ученика. Укажите столбцы вручную в диалоге импорта."
            )

        layout = HeaderLayout(header_rows=header_rows, columns=columns, student_column=student_column)
        layout.total_column = self._detect_total_column(columns)
        layout.competition_column = self._detect_competition_column(columns)

        if layout.competition_column is None:
            # Лист-матрица: под шапкой может быть ещё несколько строк подписей.
            layout.header_rows = self._extend_header_block(
                window, student_column, header_rows, max_row
            )
            if layout.header_rows != header_rows:
                layout.columns = self._build_columns(window, layout.header_rows, max_column)

        layout.criterion_columns = self._detect_criterion_columns(layout.columns, layout)
        layout.login_column = self._detect_login_column(layout.columns, layout)
        self._build_display_names(layout.columns, layout.criterion_columns)
        self._assign_fields(layout)

        if layout.competition_column is None and layout.criterion_columns:
            layout.mode = MODE_CRITERIA
        if self.mode_override:
            layout.mode = self.mode_override
        return layout

    def _build_columns(
        self, window: Sequence[tuple], header_rows: tuple[int, int], max_column: int
    ) -> dict[int, ColumnInfo]:
        """Собирает сведения о столбцах по блоку заголовков."""
        labels = self._collect_labels(
            window, header_rows[0], header_rows[1] - header_rows[0] + 1, max_column
        )
        columns: dict[int, ColumnInfo] = {}
        for column in range(1, max_column + 1):
            own, label, group = labels.get(column, ("", "", ""))
            columns[column] = ColumnInfo(
                index=column,
                label=label or own,
                group_label=group,
            )
        return columns

    @staticmethod
    def _assign_fields(layout: HeaderLayout) -> None:
        """Проставляет, за что отвечает каждый столбец."""
        for column, info in layout.columns.items():
            if column == layout.student_column:
                info.field = "student_name"
            elif column == layout.total_column:
                info.field = "total_score"
            elif column == layout.competition_column:
                info.field = "competition_name"

    @staticmethod
    def _criteria_count(labels: dict[int, tuple[str, str, str]]) -> int:
        """Сколько столбцов блока похожи на критерии (подписанные, не служебные)."""
        counter = 0
        for own, label, _group in labels.values():
            text = label or own
            if not text or is_service_label(text) or is_total_marker(text):
                continue
            if exact_match_column(text) in ("student_name", "competition_name", "total_score"):
                continue
            counter += 1
        return counter

    @staticmethod
    def _student_column_from_labels(
        labels: dict[int, tuple[str, str, str]]
    ) -> tuple[int | None, bool]:
        """Столбец с ФИО по подписям блока и признак строгого совпадения."""
        exact: int | None = None
        partial: int | None = None
        for column in sorted(labels):
            own, label, _group = labels[column]
            text = label or own
            if not text:
                continue
            if exact_match_column(text) == "student_name":
                exact = exact or column
            elif (
                partial is None
                and match_column(text) == "student_name"
                and len(text) <= MAX_STUDENT_LABEL_LENGTH
            ):
                partial = column
        if exact is not None:
            return exact, True
        return partial, False

    def _detect_student_column(self, columns: dict[int, ColumnInfo]) -> int | None:
        """Ищет столбец с ФИО ученика."""
        labels = {
            column: (info.label, info.label, info.group_label)
            for column, info in columns.items()
            if info.label
        }
        column, _exact = self._student_column_from_labels(labels)
        return column

    @staticmethod
    def _detect_total_column(columns: dict[int, ColumnInfo]) -> int | None:
        """Ищет столбец с итоговой суммой (только точное совпадение подписи).

        Точное сравнение обязательно: критерии вида «Призовые места в заочном
        конкурсе» не должны приниматься за столбец «Итого».
        """
        for column, info in sorted(columns.items()):
            if info.label and exact_match_column(info.label) == "total_score":
                return column
        return None

    @staticmethod
    def _detect_competition_column(columns: dict[int, ColumnInfo]) -> int | None:
        """Ищет столбец «Название конкурса» (только точное совпадение)."""
        for column, info in sorted(columns.items()):
            if exact_match_column(info.label) == "competition_name":
                return column
        return None

    @staticmethod
    def _detect_criterion_columns(
        columns: dict[int, ColumnInfo], layout: HeaderLayout
    ) -> tuple[int, ...]:
        """Столбцы-критерии: подписанные числовые столбцы, кроме служебных."""
        service = {layout.student_column}
        for column in (layout.total_column, layout.competition_column):
            if column is not None:
                service.add(column)

        result: list[int] = []
        for column, info in sorted(columns.items()):
            if column in service or not info.label:
                continue
            if is_service_label(info.label) or is_total_marker(info.label):
                continue
            result.append(column)
        return tuple(result)

    @staticmethod
    def _detect_login_column(
        columns: dict[int, ColumnInfo], layout: HeaderLayout
    ) -> int | None:
        """Служебный столбец с баллами входа («за вход в систему»).

        Он не является критерием, но его баллы в реальных файлах выпадают
        из формулы ИТОГО — при разборе они добавляются к итогу ученика.
        """
        for column, info in sorted(columns.items()):
            if column == layout.student_column:
                continue
            if info.label and is_login_column(info.label):
                return column
        return None

    @staticmethod
    def _build_display_names(columns: dict[int, ColumnInfo], criteria: Iterable[int]) -> None:
        """Формирует читаемые имена критериев (с группой при повторах)."""
        criteria = list(criteria)
        counts: dict[str, int] = {}
        for column in criteria:
            label = columns[column].label
            counts[label] = counts.get(label, 0) + 1

        for column in criteria:
            info = columns[column]
            if counts.get(info.label, 0) > 1 and info.group_label:
                info.display_name = f"{shorten(info.group_label)} → {info.label}"
            else:
                info.display_name = info.label

    def detect_header(
        self, worksheet: Any, header_row: int | None = None
    ) -> tuple[int, dict[str, int]]:
        """Находит строку заголовков и соответствие «поле → номер столбца».

        Вызывается при построении структуры по ручным указаниям пользователя.
        """
        if header_row is None:
            layout = self.detect_layout(worksheet)
            return layout.first_header_row, self._layout_mapping(layout)

        max_column = worksheet.max_column or 0
        values = self._read_rows(worksheet, header_row, header_row, max_column)[0]
        mapping = detect_columns(values)
        if "student_name" not in mapping:
            raise ExcelParseError(f"Строка {header_row}: {DataValidator.describe_missing(mapping)}")
        return header_row, mapping

    @staticmethod
    def _layout_mapping(layout: HeaderLayout) -> dict[str, int]:
        mapping: dict[str, int] = {"student_name": layout.student_column}
        if layout.total_column is not None:
            mapping["total_score"] = layout.total_column
        if layout.competition_column is not None:
            mapping["competition_name"] = layout.competition_column
        return mapping

    # --- Разбор листа --------------------------------------------------------

    def _layout_with_overrides(
        self, worksheet: Any, header_row: int | None, columns: dict[str, int] | None
    ) -> HeaderLayout:
        """Строит структуру с учётом ручных указаний пользователя."""
        if not header_row and not columns and not self.header_row_override and not self.column_overrides:
            return self.detect_layout(worksheet)

        if header_row is not None:
            self.header_row_override = header_row
        if columns:
            self.column_overrides = {str(k): int(v) for k, v in columns.items()}
        return self._layout_from_overrides(worksheet)

    def _layout_from_overrides(self, worksheet: Any) -> HeaderLayout:
        """Структура по ручным указаниям пользователя."""
        max_column = worksheet.max_column or 0
        max_row = worksheet.max_row or 0
        header_row = self.header_row_override or 1

        student_column = self.column_overrides.get("student_name")
        total_column = self.column_overrides.get("total_score")
        competition_column = self.column_overrides.get("competition_name")

        if not student_column or (self.mode_override is None and not total_column and not competition_column):
            _row, detected = self.detect_header(worksheet, header_row=header_row)
            student_column = student_column or detected.get("student_name")
            total_column = total_column or detected.get("total_score")
            competition_column = competition_column or detected.get("competition_name")

        if not student_column:
            raise ExcelParseError("Не указан столбец с ФИО ученика")

        layout = HeaderLayout(
            header_rows=(header_row, header_row),
            columns={},
            student_column=int(student_column),
            total_column=int(total_column) if total_column else None,
            competition_column=int(competition_column) if competition_column else None,
        )
        window = self._header_window(worksheet, max_column, max_row)
        layout.header_rows = self._extend_header_block(
            window, layout.student_column, layout.header_rows, max_row
        )
        layout.columns = self._build_columns(window, layout.header_rows, max_column)
        layout.criterion_columns = self._detect_criterion_columns(layout.columns, layout)
        layout.login_column = self._detect_login_column(layout.columns, layout)
        self._build_display_names(layout.columns, layout.criterion_columns)
        self._assign_fields(layout)

        mode = self.mode_override
        if mode is None:
            mode = MODE_ROWS if layout.competition_column is not None else MODE_CRITERIA
        layout.mode = mode
        return layout

    def parse_sheet(
        self,
        sheet_name: str | None = None,
        header_row: int | None = None,
        columns: dict[str, int] | None = None,
        mode: str | None = None,
    ) -> SheetParseResult:
        """Разбирает один лист в :class:`Direction`."""
        worksheet = self.worksheet(sheet_name)
        resolved_name = sheet_name or worksheet.title
        normalized_name = normalize_header(resolved_name)

        layout = self._layout_with_overrides(worksheet, header_row, columns)
        if mode is not None:
            layout.mode = mode
        if layout.mode == MODE_CRITERIA and not layout.criterion_columns:
            raise ExcelParseError(
                f"Лист «{resolved_name}»: не найдено ни одного столбца-критерия с баллами."
            )
        if layout.mode == MODE_ROWS and layout.total_column is None:
            raise ExcelParseError(
                f"Лист «{resolved_name}»: не найден столбец с баллами («Итого»). "
                "Укажите столбцы вручную в диалоге импорта."
            )

        sink = _WarningSink(MAX_SHEET_WARNINGS)
        skipped_rows: list[int] = []
        if layout.mode == MODE_CRITERIA:
            direction, rows_parsed = self._parse_criteria_rows(
                worksheet, layout, resolved_name, sink, skipped_rows
            )
        else:
            direction, rows_parsed = self._parse_record_rows(
                worksheet, layout, resolved_name, sink, skipped_rows
            )

        warnings = sink.messages
        if not direction.evaluations:
            warnings.append("на листе не найдено ни одной корректной строки с оценками")
        if any(marker in normalized_name for marker in SUMMARY_SHEET_MARKERS) and len(normalized_name) <= 24:
            warnings.append(
                "лист похож на общий свод: возможен двойной учёт баллов, "
                "если те же данные загружены из файлов преподавателей"
            )
        if self.external_links:
            warnings.append(
                "файл ссылается на другие книги Excel — обычно это сводный файл, "
                "а не данные отдельного преподавателя"
            )

        self.warnings.extend(f"{resolved_name}: {warning}" for warning in warnings)

        return SheetParseResult(
            sheet_name=resolved_name,
            direction=direction,
            header_row=layout.first_header_row,
            header_rows=layout.header_rows,
            columns=self._layout_mapping(layout),
            mode=layout.mode,
            criterion_columns=layout.criterion_columns,
            criterion_names=tuple(layout.display_name(column) for column in layout.criterion_columns),
            skipped_rows=skipped_rows,
            warnings=warnings,
            rows_parsed=rows_parsed,
        )

    def _iter_data_rows(self, worksheet: Any, layout: HeaderLayout):
        """Итерирует строки данных: (номер строки, значения строки).

        Итерация прекращается после длинной серии строк без ФИО: реальные листы
        содержат тысячи «фантомных» пустых строк (остатки форматирования).
        """
        max_column = worksheet.max_column or 0
        max_row = worksheet.max_row or 0
        if layout.first_data_row > max_row:
            return

        empty_run = 0
        for offset, values in enumerate(
            worksheet.iter_rows(
                min_row=layout.first_data_row,
                max_row=max_row,
                max_col=max_column,
                values_only=True,
            )
        ):
            values = tuple(values)
            if not normalize_name(self._cell(values, layout.student_column)):
                empty_run += 1
                if empty_run >= MAX_EMPTY_ROWS_RUN:
                    break
            else:
                empty_run = 0
            yield layout.first_data_row + offset, values

    @staticmethod
    def _row_is_empty(values: Sequence[Any]) -> bool:
        """Полностью пустая строка."""
        for value in values:
            if value is None:
                continue
            if isinstance(value, str) and not value.strip():
                continue
            return False
        return True

    def _parse_record_rows(
        self,
        worksheet: Any,
        layout: HeaderLayout,
        sheet_name: str,
        sink: _WarningSink,
        skipped_rows: list[int],
    ) -> tuple[Direction, int]:
        """Режим «строки»: ФИО | конкурс | итого (шаблон из плана)."""
        direction = Direction(name=sheet_name)
        rows_parsed = 0
        has_competition = layout.competition_column is not None
        has_total = layout.total_column is not None

        for row_index, values in self._iter_data_rows(worksheet, layout):
            if self._row_is_empty(values):
                continue

            student_value = self._cell(values, layout.student_column)
            competition_value = self._cell(values, layout.competition_column) if has_competition else None
            score_value = self._cell(values, layout.total_column) if has_total else None

            student_name = normalize_name(student_value)
            competition_name = normalize_name(competition_value)

            if not student_name:
                if is_total_marker(competition_name) or is_total_marker(score_value):
                    continue
                if competition_name or normalize_name(score_value):
                    skipped_rows.append(row_index)
                    sink.add(f"строка {row_index}: не указано ФИО — строка пропущена")
                continue

            if is_header_or_total_marker(student_name):
                # Повторный заголовок или строка «ИТОГО» внутри листа.
                continue

            total_score = parse_score(score_value)
            if total_score is None:
                if score_value is None or (isinstance(score_value, str) and not score_value.strip()):
                    # Пустая ячейка баллов — ноль баллов: ученик попадает в сводную таблицу
                    # (сюда же попадает формула без вычисленного кэша — она и читается как пустая).
                    total_score = 0.0
                    sink.add(f"строка {row_index}: не указано значение баллов — учтено 0")
                else:
                    skipped_rows.append(row_index)
                    sink.add(
                        f"строка {row_index}: значение баллов «{score_value}» не является числом "
                        "— строка пропущена"
                    )
                    continue

            if has_competition and not competition_name:
                sink.add(f"строка {row_index}: не указано название конкурса")

            direction.add_evaluation(
                Evaluation(
                    student_name=student_name,
                    competition_name=competition_name,
                    total_score=total_score,
                )
            )
            rows_parsed += 1

        return direction, rows_parsed

    def _parse_criteria_rows(
        self,
        worksheet: Any,
        layout: HeaderLayout,
        sheet_name: str,
        sink: _WarningSink,
        skipped_rows: list[int],
    ) -> tuple[Direction, int]:
        """Режим «критерии»: баллы начисляются в столбцах-критериях.

        Если на листе есть колонка ИТОГО, её значение считается контрольным:
        при расхождении со суммой критериев в проект идёт значение из файла,
        а расхождение попадает в предупреждения.

        Формула ИТОГО реальных файлов (``=SUM(C:AF)``) не учитывает служебный
        столбец «за вход в систему»: если файловое ИТОГО равно сумме критериев,
        баллы этого столбца прибавляются к итогу отдельной оценкой с
        предупреждением. Строка с корректным ФИО всегда даёт запись, даже
        если ученик не набрал ни одного балла (0 баллов).
        """
        direction = Direction(name=sheet_name)
        rows_parsed = 0

        for row_index, values in self._iter_data_rows(worksheet, layout):
            if self._row_is_empty(values):
                continue

            student_name = normalize_name(self._cell(values, layout.student_column))
            if not student_name or is_header_or_total_marker(student_name):
                continue

            awards: list[tuple[str, float]] = []
            for column in layout.criterion_columns:
                raw_value = self._cell(values, column)
                if raw_value is None or (isinstance(raw_value, str) and not raw_value.strip()):
                    continue
                points = parse_score(raw_value)
                if points is None:
                    sink.add(
                        f"строка {row_index}: значение «{raw_value}» в столбце "
                        f"«{shorten(layout.display_name(column), 40)}» не является числом — пропущено"
                    )
                    continue
                if points == 0:
                    continue  # ноль баллов = критерий не выполнен
                awards.append((layout.display_name(column), points))

            total_value = (
                parse_score(self._cell(values, layout.total_column))
                if layout.total_column
                else None
            )
            awards_sum = float(sum(points for _name, points in awards))

            # Баллы служебного столбца «за вход в систему»: читаются только при
            # наличии колонки ИТОГО (без неё формулы, которая могла бы их
            # пропустить, в файле нет).
            login_points = 0.0
            if total_value is not None and layout.login_column is not None:
                login_points = parse_score(self._cell(values, layout.login_column)) or 0.0
            expected_total = awards_sum + login_points
            matches_awards = total_value is not None and abs(total_value - awards_sum) <= 0.001
            matches_expected = total_value is not None and abs(total_value - expected_total) <= 0.001

            if awards and total_value is not None and not matches_awards and not matches_expected:
                if awards_sum > 0 and total_value <= 0:
                    sink.add(
                        f"строка {row_index}: колонка ИТОГО не заполнена ({total_value:g}) — "
                        f"использована сумма критериев ({awards_sum:g})"
                    )
                else:
                    sink.add(
                        f"строка {row_index}: сумма критериев ({awards_sum:g}) не совпадает "
                        f"с колонкой ИТОГО ({total_value:g}) — взято значение из файла"
                    )
                    direction.add_evaluation(
                        Evaluation(
                            student_name=student_name,
                            competition_name="Итого по листу (из файла)",
                            total_score=total_value,
                        )
                    )
                    rows_parsed += 1
                    continue

            # Файл «забыл» столбец входа в формуле ИТОГО (или, наоборот,
            # учёл его) — баллы столбца добавляются к итогу отдельной оценкой.
            login_recorded = False
            if login_points > 0 and (matches_awards or matches_expected):
                if matches_awards:
                    login_label = shorten(layout.display_name(layout.login_column), 40)
                    sink.add(
                        f"строка {row_index}: формула ИТОГО не учитывает столбец «{login_label}» "
                        f"({login_points:g} балл.) — добавлено к итогу"
                    )
                direction.add_evaluation(
                    Evaluation(
                        student_name=student_name,
                        competition_name=layout.display_name(layout.login_column),
                        total_score=login_points,
                    )
                )
                login_recorded = True

            if awards:
                for criterion_name, points in awards:
                    direction.add_evaluation(
                        Evaluation(
                            student_name=student_name,
                            competition_name=criterion_name,
                            total_score=points,
                        )
                    )
                rows_parsed += 1
            else:
                if not login_recorded:
                    # Ученик без единого балла тоже попадает в сводную таблицу.
                    direction.add_evaluation(
                        Evaluation(
                            student_name=student_name,
                            competition_name="",
                            total_score=total_value or 0.0,
                        )
                    )
                rows_parsed += 1

        return direction, rows_parsed

    # --- Разбор книги --------------------------------------------------------

    def parse_all(
        self,
        sheet_names: list[str] | None = None,
        strict: bool = False,
    ) -> dict[str, SheetParseResult]:
        """Разбирает все (или указанные) листы книги.

        Ошибка на одном листе не прерывает разбор: лист получает пустое
        направление и предупреждение (если ``strict`` не задан).
        Используется только в тестах — интерфейс разбирает листы поштучно
        через :meth:`parse_sheet` (см. :func:`collect_upload_items`).
        """
        names = list(sheet_names) if sheet_names else self.sheet_names
        results: dict[str, SheetParseResult] = {}
        for name in names:
            try:
                results[name] = self.parse_sheet(name)
            except ExcelParseError as exc:
                if strict:
                    raise
                message = str(exc)
                self.warnings.append(message)
                results[name] = SheetParseResult(
                    sheet_name=name,
                    direction=Direction(name=name),
                    warnings=[message],
                )
        return results

    def parse_all_teachers(
        self, sheet_to_teacher: dict[str, str | None] | None = None
    ) -> list[Teacher]:
        """Разбирает книгу и группирует листы по преподавателям.

        :param sheet_to_teacher: необязательный маппинг «лист → ФИО преподавателя».
            Значение ``None`` означает «лист не импортировать».

        Используется только в тестах: интерфейс сопоставляет листы
        преподавателям при импорте сам (см. :func:`collect_upload_items`).
        """
        sheet_names = list(sheet_to_teacher) if sheet_to_teacher else None
        results = self.parse_all(sheet_names=sheet_names)

        teachers: list[Teacher] = []
        by_key: dict[str, Teacher] = {}

        for sheet_name, result in results.items():
            mapping: str | None = None
            if sheet_to_teacher is not None and sheet_name in sheet_to_teacher:
                mapping = sheet_to_teacher[sheet_name]
                if mapping is None:
                    continue

            teacher_name = normalize_name(mapping or "") or self.default_teacher_name() or sheet_name
            key = name_key(teacher_name)
            teacher = by_key.get(key)
            if teacher is None:
                teacher = Teacher(name=teacher_name)
                by_key[key] = teacher
                teachers.append(teacher)

            if teacher.get_direction(result.direction.name) is None:
                teacher.add_direction(result.direction, replace=False)

        return teachers

class _WarningSink:
    """Собирает предупреждения по листу, ограничивая их количество."""

    def __init__(self, limit: int) -> None:
        self.limit = max(1, limit)
        self.messages: list[str] = []
        self.total = 0

    def add(self, message: str) -> None:
        """Добавляет предупреждение (лишние считаются, но не показываются)."""
        self.total += 1
        if len(self.messages) < self.limit:
            self.messages.append(message)
        elif len(self.messages) == self.limit:
            self.messages.append("…остальные предупреждения по листу скрыты")
