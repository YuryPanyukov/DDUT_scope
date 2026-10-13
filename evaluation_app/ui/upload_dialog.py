"""Диалоги загрузки Excel: предпросмотр и ручное определение столбцов (п. 6.1, 8.2)."""

from __future__ import annotations

from collections.abc import Sequence
from dataclasses import dataclass, field

from openpyxl.utils import get_column_letter
from PyQt5.QtCore import Qt, pyqtSignal
from PyQt5.QtWidgets import (
    QAbstractItemView,
    QCheckBox,
    QComboBox,
    QDialog,
    QDialogButtonBox,
    QFormLayout,
    QHBoxLayout,
    QHeaderView,
    QLabel,
    QPlainTextEdit,
    QPushButton,
    QSpinBox,
    QTableWidget,
    QTableWidgetItem,
    QVBoxLayout,
    QWidget,
)

from evaluation_app import config
from evaluation_app.services.excel_parser import (
    MAX_STUDENT_LABEL_LENGTH,
    MODE_CRITERIA,
    MODE_ROWS,
    ExcelParseError,
    ExcelParser,
    SheetParseResult,
)
from evaluation_app.ui.column_controls import enable_column_controls
from evaluation_app.utils.file_utils import file_name, format_score
from evaluation_app.utils.validators import (
    FIELD_TITLES,
    detect_columns_exact,
    match_column,
    missing_fields,
    normalize_name,
)

SKIP_TEACHER_LABEL = "<пропустить лист>"


@dataclass
class UploadSheet:
    """Один лист Excel, подготовленный к импорту."""

    file_path: str
    sheet_name: str
    parser: ExcelParser | None = None
    default_teacher: str = ""
    result: SheetParseResult | None = None
    error: str | None = None
    warnings: list[str] = field(default_factory=list)

    @property
    def rows_count(self) -> int:
        """Количество разобранных строк."""
        return self.result.evaluations_count if self.result else 0

    @property
    def total_score(self) -> float:
        """Сумма баллов листа."""
        return self.result.total_score if self.result else 0.0

    def all_warnings(self) -> list[str]:
        """Все предупреждения листа."""
        messages = list(self.warnings)
        if self.result is not None:
            messages.extend(self.result.warnings)
        if self.error:
            messages.append(self.error)
        return messages


class UploadPreviewDialog(QDialog):
    """Предпросмотр загружаемых файлов и привязка листов к преподавателям."""

    configureRequested = pyqtSignal(int)  # индекс строки таблицы

    def __init__(
        self,
        items: Sequence[UploadSheet],
        existing_teachers: Sequence[str] = (),
        parent: QWidget | None = None,
    ) -> None:
        super().__init__(parent)
        self.setWindowTitle("Загрузка Excel-файлов")
        self.resize(980, 700)
        self._items = list(items)
        self._existing_teachers = [name for name in existing_teachers if name]

        layout = QVBoxLayout(self)

        hint = QLabel(
            "Проверьте, какие листы и к какому преподавателю будут привязаны. "
            "Листы, которые не нужно импортировать, отметьте как «пропустить».",
            self,
        )
        hint.setWordWrap(True)
        layout.addWidget(hint)

        self.mapping_table = QTableWidget(self)
        self.mapping_table.setColumnCount(6)
        self.mapping_table.setHorizontalHeaderLabels(
            ["Файл", "Лист", "Преподаватель", "Строк", "Сумма баллов", "Столбцы"]
        )
        self.mapping_table.setAlternatingRowColors(True)
        self.mapping_table.setSelectionBehavior(QAbstractItemView.SelectRows)
        self.mapping_table.setEditTriggers(QAbstractItemView.NoEditTriggers)
        self.mapping_table.verticalHeader().setVisible(False)
        header = self.mapping_table.horizontalHeader()
        header.setSectionResizeMode(0, QHeaderView.ResizeToContents)
        header.setSectionResizeMode(1, QHeaderView.ResizeToContents)
        header.setSectionResizeMode(2, QHeaderView.ResizeToContents)
        header.setSectionResizeMode(3, QHeaderView.ResizeToContents)
        header.setSectionResizeMode(4, QHeaderView.ResizeToContents)
        header.setSectionResizeMode(5, QHeaderView.Stretch)
        self.mapping_table.currentCellChanged.connect(lambda *_: self._update_preview())
        layout.addWidget(self.mapping_table, 3)
        enable_column_controls(self.mapping_table)

        preview_row = QHBoxLayout()
        preview_row.addWidget(QLabel("Предпросмотр данных:", self))
        preview_row.addStretch(1)
        self.configure_button = QPushButton("Указать столбцы вручную...", self)
        self.configure_button.clicked.connect(self._on_configure)
        preview_row.addWidget(self.configure_button)
        layout.addLayout(preview_row)

        self.preview_table = QTableWidget(self)
        self.preview_table.setColumnCount(3)
        self.preview_table.setHorizontalHeaderLabels(["ФИО", "Конкурс / критерий", "Баллы"])
        self.preview_table.setEditTriggers(QAbstractItemView.NoEditTriggers)
        self.preview_table.setAlternatingRowColors(True)
        self.preview_table.verticalHeader().setVisible(False)
        preview_header = self.preview_table.horizontalHeader()
        preview_header.setSectionResizeMode(0, QHeaderView.Stretch)
        preview_header.setSectionResizeMode(1, QHeaderView.Stretch)
        preview_header.setSectionResizeMode(2, QHeaderView.ResizeToContents)
        layout.addWidget(self.preview_table, 3)
        enable_column_controls(self.preview_table)

        layout.addWidget(QLabel("Предупреждения:", self))
        self.warnings_view = QPlainTextEdit(self)
        self.warnings_view.setReadOnly(True)
        self.warnings_view.setMaximumHeight(120)
        layout.addWidget(self.warnings_view)

        self.buttons = QDialogButtonBox(QDialogButtonBox.Ok | QDialogButtonBox.Cancel, self)
        self.buttons.button(QDialogButtonBox.Ok).setText("Импортировать")
        self.buttons.button(QDialogButtonBox.Cancel).setText("Отмена")
        self.buttons.accepted.connect(self.accept)
        self.buttons.rejected.connect(self.reject)
        layout.addWidget(self.buttons)

        self.refresh()

    # --- API -----------------------------------------------------------------

    def refresh(self) -> None:
        """Перерисовывает таблицу привязок, предпросмотр и предупреждения."""
        self.mapping_table.blockSignals(True)
        self.mapping_table.setRowCount(len(self._items))

        for row, item in enumerate(self._items):
            file_item = QTableWidgetItem(file_name(item.file_path))
            file_item.setToolTip(item.file_path)
            self.mapping_table.setItem(row, 0, file_item)
            self.mapping_table.setItem(row, 1, QTableWidgetItem(item.sheet_name))

            combo = QComboBox(self)
            combo.setEditable(True)
            combo.addItem(SKIP_TEACHER_LABEL, None)
            for name in self._teacher_candidates(item):
                combo.addItem(name, name)
            index = combo.findData(item.default_teacher)
            combo.setCurrentIndex(index if index >= 0 else (1 if combo.count() > 1 else 0))
            combo.currentIndexChanged.connect(self._update_import_button)
            combo.editTextChanged.connect(self._update_import_button)
            self.mapping_table.setCellWidget(row, 2, combo)

            rows_item = QTableWidgetItem(str(item.rows_count))
            rows_item.setTextAlignment(Qt.AlignRight | Qt.AlignVCenter)
            self.mapping_table.setItem(row, 3, rows_item)

            score_item = QTableWidgetItem(format_score(item.total_score))
            score_item.setTextAlignment(Qt.AlignRight | Qt.AlignVCenter)
            self.mapping_table.setItem(row, 4, score_item)

            columns_text = "структура не определена"
            if item.result and item.result.columns:
                columns_text = item.result.describe_columns()
            columns_cell = QTableWidgetItem(columns_text)
            crit = len(item.result.criterion_columns) if item.result else 0
            columns_cell.setToolTip(
                "Лист-матрица: баллы начисляются по критериям"
                if crit
                else "Обычная таблица: одна строка — одна оценка"
            )
            self.mapping_table.setItem(row, 5, columns_cell)

        self.mapping_table.blockSignals(False)
        if self._items and self.mapping_table.currentRow() < 0:
            self.mapping_table.setCurrentCell(0, 0)

        self._update_preview()
        self._update_warnings()
        self._update_import_button()

    def get_selection(self) -> list[tuple[UploadSheet, str | None]]:
        """Выбранные для импорта листы: (элемент предпросмотра, ФИО преподавателя).

        Возвращается именно элемент :class:`UploadSheet`, а не путь к файлу:
        лист уже разобран (в том числе с ручными настройками столбцов), и
        повторное чтение файла при импорте не нужно.
        """
        selection: list[tuple[UploadSheet, str | None]] = []
        for row, item in enumerate(self._items):
            teacher_name = self._teacher_for_row(row)
            if teacher_name is None:
                continue
            selection.append((item, teacher_name))
        return selection

    def item_at(self, row: int) -> UploadSheet | None:
        """Элемент предпросмотра по индексу строки."""
        if 0 <= row < len(self._items):
            return self._items[row]
        return None

    # --- Внутреннее ----------------------------------------------------------

    def _teacher_candidates(self, item: UploadSheet) -> list[str]:
        candidates: list[str] = []
        for name in [*self._existing_teachers, item.default_teacher]:
            if name and name not in candidates:
                candidates.append(name)
        return candidates

    def _teacher_for_row(self, row: int) -> str | None:
        widget = self.mapping_table.cellWidget(row, 2)
        if not isinstance(widget, QComboBox):
            return None
        text = widget.currentText().strip()
        if not text or text == SKIP_TEACHER_LABEL:
            return None
        return text

    def _update_preview(self) -> None:
        row = self.mapping_table.currentRow()
        self.preview_table.setRowCount(0)
        self.preview_table.setHorizontalHeaderLabels(["ФИО", "Конкурс / критерий", "Баллы"])
        if row < 0 or row >= len(self._items):
            return

        item = self._items[row]
        rows = item.result.to_rows(limit=config.PREVIEW_ROWS_LIMIT) if item.result else []
        self.preview_table.setRowCount(len(rows))
        for index, values in enumerate(rows):
            for column, value in enumerate(values[:3]):
                cell = QTableWidgetItem(str(value))
                if column == 2:
                    cell.setTextAlignment(Qt.AlignRight | Qt.AlignVCenter)
                self.preview_table.setItem(index, column, cell)

    def _update_warnings(self) -> None:
        lines: list[str] = []
        for item in self._items:
            for warning in item.all_warnings():
                lines.append(f"{file_name(item.file_path)} / {item.sheet_name}: {warning}")
        self.warnings_view.setPlainText("\n".join(lines) if lines else "Предупреждений нет.")
        self._update_import_button()

    def _update_import_button(self) -> None:
        selection = self.get_selection()
        chosen = {(item.file_path, item.sheet_name) for item, _teacher in selection}
        rows_total = sum(
            item.rows_count
            for item in self._items
            if (item.file_path, item.sheet_name) in chosen
        )
        ok_button = self.buttons.button(QDialogButtonBox.Ok)
        ok_button.setText(f"Импортировать ({len(selection)} лист(ов), {rows_total} строк)")
        ok_button.setEnabled(bool(selection))
        self.configure_button.setEnabled(self.mapping_table.currentRow() >= 0)

    def _on_configure(self) -> None:
        row = self.mapping_table.currentRow()
        if row >= 0:
            self.configureRequested.emit(row)


class ColumnMappingDialog(QDialog):
    """Ручное указание строки заголовков и столбцов (п. 7 плана)."""

    def __init__(self, parser: ExcelParser, sheet_name: str, parent: QWidget | None = None) -> None:
        super().__init__(parent)
        self.setWindowTitle(f"Настройка столбцов — {sheet_name}")
        self.resize(760, 520)

        self._rows = parser.preview_rows(sheet_name)
        self._header_row = 1
        self._columns: dict[str, int] = {}

        layout = QVBoxLayout(self)
        layout.addWidget(
            QLabel(
                "Выберите строку с заголовками и укажите, какой столбец за что отвечает.",
                self,
            )
        )

        self.grid = QTableWidget(self)
        self.grid.setColumnCount(max((len(row) for row in self._rows), default=0))
        self.grid.setRowCount(len(self._rows))
        self.grid.setHorizontalHeaderLabels(
            [get_column_letter(col) for col in range(1, self.grid.columnCount() + 1)]
        )
        self.grid.verticalHeader().setVisible(True)
        self.grid.setEditTriggers(QAbstractItemView.NoEditTriggers)
        self.grid.setSelectionBehavior(QAbstractItemView.SelectRows)
        for row_index, values in enumerate(self._rows, start=1):
            self.grid.setVerticalHeaderItem(row_index - 1, QTableWidgetItem(str(row_index)))
            for column_index, value in enumerate(values):
                text = "" if value is None else str(value)
                self.grid.setItem(row_index - 1, column_index, QTableWidgetItem(text))
        layout.addWidget(self.grid, 1)
        enable_column_controls(self.grid)

        form = QFormLayout()
        self.header_spin = QSpinBox(self)
        self.header_spin.setMinimum(1)
        self.header_spin.setMaximum(max(len(self._rows), 1))
        self.header_spin.valueChanged.connect(self._on_header_row_changed)
        form.addRow("Строка заголовков:", self.header_spin)

        self.field_combos: dict[str, QComboBox] = {}
        for field_name, title in FIELD_TITLES.items():
            combo = QComboBox(self)
            form.addRow(f"{title}:", combo)
            self.field_combos[field_name] = combo

        self.criteria_checkbox = QCheckBox(
            "Лист-матрица: считать столбцы с подписями критериями (баллы начисляются в ячейках)",
            self,
        )
        self.criteria_checkbox.setToolTip(
            "Отметьте для оценочных листов вида «ФИО | критерий 1 | ... | критерий N | ИТОГО».\n"
            "В этом режиме обязателен только столбец с ФИО, а колонка ИТОГО — необязательна."
        )
        self.criteria_checkbox.toggled.connect(self._on_mode_toggled)
        form.addRow("Режим разбора:", self.criteria_checkbox)

        layout.addLayout(form)
        self.status_label = QLabel("", self)
        self.status_label.setWordWrap(True)
        layout.addWidget(self.status_label)

        buttons = QDialogButtonBox(QDialogButtonBox.Ok | QDialogButtonBox.Cancel, self)
        buttons.button(QDialogButtonBox.Ok).setText("Применить")
        buttons.button(QDialogButtonBox.Cancel).setText("Отмена")
        buttons.accepted.connect(self._on_accept)
        buttons.rejected.connect(self.reject)
        layout.addWidget(buttons)

        # Попробуем определить структуру (строку заголовков и режим) автоматически.
        try:
            layout = parser.detect_layout(parser.worksheet(sheet_name))
            self.header_spin.setValue(layout.first_header_row)
            self.criteria_checkbox.setChecked(layout.mode == MODE_CRITERIA)
        except ExcelParseError:
            self.header_spin.setValue(1)
        self._on_header_row_changed(self.header_spin.value())

    # --- API -----------------------------------------------------------------

    def header_row(self) -> int:
        """Выбранная строка заголовков."""
        return self._header_row

    def columns(self) -> dict[str, int]:
        """Выбранные столбцы «поле → номер (с 1)»."""
        return dict(self._columns)

    def mode(self) -> str:
        """Режим разбора: «rows» (строка = оценка) или «criteria» (лист-матрица)."""
        return MODE_CRITERIA if self.criteria_checkbox.isChecked() else MODE_ROWS

    # --- Внутреннее ----------------------------------------------------------

    def _on_mode_toggled(self, checked: bool) -> None:
        """В режиме критериев обязателен только столбец с ФИО."""
        self._validate()

    def _on_header_row_changed(self, row: int) -> None:
        self._header_row = int(row)
        header_values = self._rows[row - 1] if 0 < row <= len(self._rows) else []
        auto_mapping = self._auto_mapping(header_values)

        for field_name, combo in self.field_combos.items():
            combo.blockSignals(True)
            combo.clear()
            combo.addItem("— не выбрано —", None)
            for column_index, value in enumerate(header_values, start=1):
                text = f"{get_column_letter(column_index)}: {value if value not in (None, '') else '—'}"
                combo.addItem(text, column_index)
            target = auto_mapping.get(field_name)
            if target is None:
                target = self._current_selection(field_name)
            index = combo.findData(target)
            combo.setCurrentIndex(index if index >= 0 else 0)
            combo.blockSignals(False)
            combo.currentIndexChanged.connect(lambda *_: self._validate())

        self._validate()

    @staticmethod
    def _auto_mapping(headers: Sequence) -> dict[str, int]:
        """Подбирает столбцы для предзаполнения диалога.

        Сначала только точные синонимы (чтобы описания критериев не «угадывались»),
        и лишь для ФИО допускается вхождение — например «Ученик» или
        «Фамилия Имя Отчество».
        """
        mapping = detect_columns_exact(headers)
        if "student_name" in mapping:
            return mapping

        for index, value in enumerate(headers, start=1):
            text = normalize_name(value)
            if not text or len(text) > MAX_STUDENT_LABEL_LENGTH:
                continue
            if match_column(text) == "student_name":
                mapping["student_name"] = index
                break
        return mapping

    def _current_selection(self, field_name: str) -> int | None:
        combo = self.field_combos[field_name]
        data = combo.currentData()
        return int(data) if isinstance(data, int) else None

    def _validate(self) -> None:
        self._columns = {}
        for field_name, combo in self.field_combos.items():
            data = combo.currentData()
            if isinstance(data, int):
                self._columns[field_name] = int(data)

        criteria_mode = self.criteria_checkbox.isChecked()
        missing = ["student_name"] if criteria_mode else missing_fields(self._columns)
        missing = [name for name in missing if name not in self._columns]
        ok_button = self.findChild(QDialogButtonBox)
        if missing:
            titles = ", ".join(FIELD_TITLES.get(name, name) for name in missing)
            self.status_label.setText(f"Не выбраны обязательные поля: {titles}")
            self.status_label.setStyleSheet("color: #b00;")
        else:
            text = "Структура определена: " + ", ".join(
                f"{FIELD_TITLES.get(name, name)} → {get_column_letter(col)}"
                for name, col in self._columns.items()
            )
            if criteria_mode:
                text += " (остальные подписанные столбцы будут считаться критериями)"
            self.status_label.setText(text)
            self.status_label.setStyleSheet("color: #060;")
        if ok_button is not None:
            ok_button.button(QDialogButtonBox.Ok).setEnabled(not missing)

    def _on_accept(self) -> None:
        self._validate()
        if missing_fields(self._columns) and "student_name" not in self._columns:
            return
        if self.mode() == MODE_ROWS and missing_fields(self._columns):
            return
        self.accept()
